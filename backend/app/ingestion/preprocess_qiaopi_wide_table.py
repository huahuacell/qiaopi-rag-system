from __future__ import annotations

import argparse
import json
import re
from pathlib import Path
from typing import Iterable

import pandas as pd

from app.nlp.text_normalizer import (
    normalize_newlines as shared_normalize_newlines,
    normalize_qiaopi_text,
)


BACKEND_DIR = Path(__file__).resolve().parents[2]
DEFAULT_INPUT = BACKEND_DIR / "data" / "raw" / "qiaopi_213_text.xlsx"
DEFAULT_OUTPUT_DIR = BACKEND_DIR / "data" / "processed"
DEFAULT_SHEET = "OCR_Source_Text"


WIDE_COLUMNS = [
    "record_id",
    "source_index",
    "title_reference",
    "page_reference",
    "has_full_text",
    "text_quality_level",
    "review_needed",
    "body_raw",
    "body_clean",
    "body_without_metadata",
    "body_core",
    "body_length",
    "line_count",
    "sentence_count",
    "sender",
    "recipient",
    "sender_role",
    "recipient_role",
    "relationship_type",
    "kinship_terms",
    "family_members_mentioned",
    "recipient_honorific",
    "self_reference",
    "origin_place",
    "destination_place",
    "overseas_place",
    "hometown_place",
    "place_mentions",
    "route_text",
    "is_overseas_to_hometown",
    "country_or_region",
    "date_text",
    "written_date_text",
    "year_raw",
    "year_normalized",
    "month_text",
    "day_text",
    "lunar_or_solar",
    "era_text",
    "date_confidence",
    "remittance_raw",
    "remittance_amount_text",
    "remittance_amount_number",
    "currency",
    "remittance_statement",
    "remittance_purpose",
    "receipt_instruction",
    "has_remittance",
    "opening_salutation",
    "opening_formula",
    "opening_greeting",
    "main_message",
    "safety_report",
    "family_care_statement",
    "instruction_statement",
    "closing_greeting",
    "closing_formula",
    "date_signature_line",
    "signature",
    "postscript",
    "theme_family_affection",
    "theme_remittance",
    "theme_safety",
    "theme_instruction",
    "theme_study",
    "theme_health",
    "theme_work",
    "theme_debt",
    "theme_marriage",
    "theme_funeral",
    "theme_home_building",
    "main_intent",
    "theme_tags",
    "qiaopi_opening_pattern",
    "qiaopi_closing_pattern",
    "honorific_phrases",
    "humble_phrases",
    "classical_words",
    "common_formulae",
    "remittance_formulae",
    "safety_formulae",
    "instruction_formulae",
    "style_keywords",
    "style_strength",
    "style_reference_text",
    "evidence_opening",
    "evidence_safety",
    "evidence_remittance",
    "evidence_family_care",
    "evidence_instruction",
    "evidence_closing",
    "evidence_entities",
    "retrieval_text",
    "rag_summary_text",
    "title_date_text",
    "title_sender_text",
    "title_recipient_text",
    "body_core_length",
    "ocr_uncertain_count",
    "sender_name_clean",
    "recipient_name_clean",
    "primary_kinship",
    "relation_confidence",
    "place_mentions_normalized",
    "remittance_mentions_json",
    "retrieval_keywords",
]

NEW_WIDE_COLUMNS = WIDE_COLUMNS[-12:]


METADATA_LABELS = {
    "寄批人": "sender",
    "收批人": "recipient",
    "批款": "remittance_raw",
    "日期": "date_text",
}

KINSHIP_TERMS = [
    "慈亲",
    "严亲",
    "双亲",
    "母亲",
    "父亲",
    "祖父",
    "祖母",
    "外祖父",
    "外祖母",
    "外祖父母",
    "祖慈",
    "家祖",
    "岳父",
    "岳母",
    "岳祖母",
    "岳慈亲",
    "婶母",
    "伯母",
    "叔父",
    "嫂",
    "大嫂",
    "阿嬷",
    "兄",
    "弟",
    "妻",
    "荆妻",
    "贤妻",
    "内妻",
    "内助",
    "儿",
    "男",
    "孙",
    "外孙",
    "侄",
    "媳",
    "媳妇",
    "姊",
    "姐",
    "妹",
    "叔",
    "伯",
    "婶",
    "姑",
    "妗",
    "内侄",
]

HONORIFIC_TERMS = [
    "大人",
    "尊前",
    "膝下",
    "台鉴",
    "英鉴",
    "钧鉴",
    "清鉴",
    "尊鉴",
    "赐鉴",
    "少爷",
    "先生",
    "贤妻",
    "内助",
    "收知",
    "金安",
    "福安",
]

SELF_REFERENCE_TERMS = [
    "小儿",
    "儿",
    "男",
    "愚弟",
    "愚姆",
    "愚",
    "弟",
    "侄",
    "孙",
    "外孙",
    "媳",
    "婿",
    "夫",
    "予",
    "俺",
]

PLACE_ALIASES = {
    "星洲": "新加坡",
    "星加坡": "新加坡",
    "新嘉坡": "新加坡",
    "新加坡": "新加坡",
    "叻坡": "新加坡",
    "石叻": "新加坡",
    "助坡": "新加坡",
    "暹罗": "泰国",
    "泰国": "泰国",
    "安南": "越南",
    "越南": "越南",
    "日本": "日本",
    "彭亨": "马来西亚彭亨",
    "嘭": "马来西亚彭亨",
    "槟城": "马来西亚槟城",
    "庇能": "马来西亚槟城",
    "泗水": "印尼泗水",
    "吧城": "印尼雅加达",
    "仰光": "缅甸仰光",
    "香港": "香港",
    "潮安": "广东潮安",
    "澄海": "广东澄海",
    "澄邑": "广东澄海",
    "澄蓬": "广东澄海",
    "澄莲": "广东澄海",
    "澄吧": "广东澄海",
    "揭阳": "广东揭阳",
    "饶邑": "广东饶平",
    "饒邑": "广东饶平",
    "饶隆": "广东饶平",
    "饶色": "广东饶平",
    "潮州": "广东潮州",
    "潮汕": "广东侨乡",
    "潮油": "广东侨乡",
    "汀邑": "广东澄海",
    "湾邑": "广东澄海",
    "溢邑": "广东澄海",
    "洒邑": "广东澄海",
    "汕头": "广东汕头",
    "海邑": "广东澄海",
    "东凤": "广东潮安东凤",
    "大寨": "广东潮安大寨",
    "西洋乡": "广东侨乡西洋乡",
    "广东": "广东",
}

PLACE_PREFIX_ALIASES = {
    "叻": "新加坡",
    "肋": "新加坡",
    "吓": "新加坡",
    "暹": "泰国",
    "越": "越南",
    "香": "香港",
    "外洋": "海外",
}

OVERSEAS_NORMALIZED = {
    "新加坡",
    "泰国",
    "越南",
    "日本",
    "马来西亚彭亨",
    "马来西亚槟城",
    "印尼泗水",
    "印尼雅加达",
    "缅甸仰光",
    "香港",
    "海外",
}

COUNTRY_OR_REGION_BY_PLACE = {
    "新加坡": "新加坡",
    "泰国": "泰国",
    "越南": "越南",
    "日本": "日本",
    "马来西亚彭亨": "马来西亚",
    "马来西亚槟城": "马来西亚",
    "印尼泗水": "印尼",
    "印尼雅加达": "印尼",
    "缅甸仰光": "缅甸",
    "香港": "香港",
    "海外": "海外",
    "广东潮安": "广东侨乡",
    "广东澄海": "广东侨乡",
    "广东揭阳": "广东侨乡",
    "广东饶平": "广东侨乡",
    "广东潮州": "广东侨乡",
    "广东侨乡": "广东侨乡",
    "广东汕头": "广东侨乡",
    "广东潮安东凤": "广东侨乡",
    "广东潮安大寨": "广东侨乡",
    "广东侨乡西洋乡": "广东侨乡",
    "广东": "广东侨乡",
}

MONEY_CURRENCY_ALIASES = {
    "大洋银": "大洋",
    "大洋": "大洋",
    "洋银": "银元",
    "荷银": "荷银",
    "英洋": "英洋",
    "港币": "港币",
    "国币": "国币",
    "中央币": "国币",
    "中央法币": "国币",
    "中央国币": "国币",
    "法币": "国币",
    "国洋": "国币",
    "金圆券": "金圆券",
    "金员（圆）券": "金圆券",
    "金员(圆)券": "金圆券",
    "中正银": "银元",
    "光洋银": "银元",
    "银元": "银元",
    "大银": "银元",
    "大龙银": "银元",
    "大龍银": "银元",
    "大良（银）": "银元",
    "大良(银)": "银元",
    "良（银）": "银元",
    "良(银)": "银元",
    "良银": "银元",
    "银": "银元",
    "元": "元",
}

CHINESE_DIGITS = {
    "零": 0,
    "〇": 0,
    "一": 1,
    "壹": 1,
    "二": 2,
    "贰": 2,
    "貳": 2,
    "弍": 2,
    "式": 2,
    "两": 2,
    "兩": 2,
    "三": 3,
    "叁": 3,
    "參": 3,
    "四": 4,
    "肆": 4,
    "五": 5,
    "伍": 5,
    "六": 6,
    "陆": 6,
    "陸": 6,
    "七": 7,
    "柒": 7,
    "八": 8,
    "捌": 8,
    "九": 9,
    "玖": 9,
    "乙": 1,
}

SMALL_UNITS = {"十": 10, "拾": 10, "百": 100, "佰": 100, "千": 1000, "仟": 1000}
LARGE_UNITS = {"万": 10000, "萬": 10000}
STEMS_BRANCHES = "甲乙丙丁戊己庚辛壬癸子丑寅卯辰巳午未申酉戌亥"

OPENING_PATTERNS = ["敬禀者", "叩禀者", "启者", "禀者", "尊前", "膝下", "台鉴", "收知"]
CLOSING_PATTERNS = ["谨禀", "谨上", "叩上", "谨叩", "手书", "拜上", "启", "托", "此候", "顺候", "并请", "此问", "此致"]
HUMBLE_PATTERNS = ["伏乞", "愚", "小儿", "儿", "弟", "叩", "谨", "予", "俺"]
CLASSICAL_WORDS = ["兹", "祈", "谅", "荷", "奉", "承", "俾", "惟", "盖", "甚慰", "不胜", "诸务", "切勿", "毋须", "容后", "余言后告", "如命"]
COMMON_FORMULAE = [
    "敬禀者",
    "叩禀者",
    "膝下",
    "尊前",
    "伏乞",
    "勿念",
    "兹托",
    "祈查收",
    "至祈查收",
    "谨禀",
    "叩上",
    "务望",
    "查收",
    "此候",
    "顺候",
    "金安",
    "福安",
    "不一",
    "容后",
]
REMITTANCE_FORMULAE = ["兹轮便", "寄上", "汇上", "奉上", "付去", "寄去", "查收", "如字查收", "分抹", "收讫"]
SAFETY_FORMULAE = ["平安", "勿念", "安好", "无恙", "清泰", "起居", "痊安"]
INSTRUCTION_FORMULAE = ["务望", "切勿", "须", "祈为", "嘱", "不可", "勤俭", "持家", "教训"]
MONEY_MENTION_PATTERN = re.compile(
    r"(?P<currency>中央国币|中央法币|中央币|大洋银|大洋|光洋银|中正银|洋银|荷银|英洋|港币|国币|法币|国洋|银元|金圆券|金员[（(]?圆[）)]?券|大龙银|大龍银|大银|大?良[（(]?银[）)]?|银)?"
    r"\s*[（(]?币?[）)]?\s*"
    r"(?P<amount>[零〇一壹乙二贰貳弍式两兩三叁參四肆五伍六陆陸七柒八捌九玖十拾百佰千仟万萬廿卅\d]+)"
    r"\s*(?P<unit>万元|萬元|元|圆|圓|角|毫)"
)
METADATA_LINE_PATTERN = re.compile(r"^(寄批人|收批人|批款|日期|时间|收款人|回批人)\s*[:：]")
SALUTATION_PATTERN = re.compile(r"(尊前|膝下|大人|收悉|如晤|台鉴|英鉴|钧鉴|清鉴|尊鉴|赐鉴|收知|侍右|尊右)")
SIGNATURE_FORMULA_PATTERN = re.compile(
    r"(谨禀|谨叩|谨上|叩上|手书|手泐|拜上|禀上|敬上|上言|拜启|具禀|泐|顿|草|字|启|托|缄|禀|上|手)$"
)
SIGNATURE_BODY_FEATURE_PATTERN = re.compile(r"(敬禀者|叩禀者|启者|兹|查收|寄上|汇上|付去|寄去|奉上|批款)")
DATE_LINE_PATTERN = re.compile(
    rf"^[{STEMS_BRANCHES}民国民國夏历夏曆旧历旧曆阴历陰曆农历農曆西历西曆零〇一壹乙二贰貳弍式两兩三叁參四肆五伍六陆陸七柒八捌九玖十拾百佰千仟万萬廿卅正元冬腊年月日号號初阳完桂闰\[\]【】〔〕,\s，、.。()（）/-]+(?:书|泐|寄|缄|禀|上|手)?$"
)

THEME_KEYWORDS = {
    "theme_family_affection": ["念", "双亲", "母", "父", "妻", "儿", "家中", "弟妹", "保重", "合家"],
    "theme_remittance": ["寄上", "汇上", "批款", "大洋", "银", "元", "查收", "收讫", "付去", "寄去"],
    "theme_safety": ["平安", "勿念", "安好", "无恙", "清泰", "痊安"],
    "theme_instruction": ["务望", "嘱", "祈为", "不可", "切勿", "须", "教训", "持家", "勤俭"],
    "theme_study": ["读书", "勤学", "升学", "学费", "书馆", "读"],
    "theme_health": ["保重", "病", "疾", "身体", "痊安", "安康", "强壮"],
    "theme_work": ["生意", "辞事", "职", "做工", "东家", "头家", "营生"],
    "theme_debt": ["借", "债", "欠", "赎", "典", "还", "银项"],
    "theme_marriage": ["婚", "嫁", "娶", "妻", "媳", "内助"],
    "theme_funeral": ["孝服", "丧", "亡", "葬", "忌"],
    "theme_home_building": ["厝", "屋", "建", "修", "赎厝", "家屋", "公厝"],
}

INTENT_PRIORITY = [
    ("theme_instruction", "instruction"),
    ("theme_health", "health_care"),
    ("theme_study", "study"),
    ("theme_family_affection", "family_affection"),
    ("theme_safety", "safety_report"),
    ("theme_work", "work"),
    ("theme_debt", "debt"),
    ("theme_marriage", "marriage"),
    ("theme_funeral", "funeral"),
    ("theme_home_building", "home_building"),
    ("theme_remittance", "remittance"),
]

MAIN_INTENT_BODY_KEYWORDS = {
    "theme_instruction": ["务望", "切勿", "须", "祈为", "嘱", "不可", "教训", "持家", "勤俭"],
    "theme_health": ["保重", "病", "疾", "身体", "痊安", "安康", "康强", "强壮"],
    "theme_study": ["读书", "勤学", "升学", "学费", "成绩", "英文", "学习"],
    "theme_family_affection": ["挂念", "远念", "勿念", "无烦远念", "合家", "甚念", "亲近之情", "为慰"],
    "theme_safety": ["平安", "安好", "无恙", "清泰"],
    "theme_work": ["生意", "辞事", "职", "做工", "东家", "头家", "营生", "薪金"],
    "theme_debt": ["借", "债", "欠", "赎", "典", "还", "银项"],
    "theme_marriage": ["婚", "嫁", "娶", "媳", "内助"],
    "theme_funeral": ["孝服", "丧", "亡", "葬", "忌"],
    "theme_home_building": ["厝", "屋", "建", "修", "赎厝", "家屋", "公厝"],
}


def safe_text(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value)


def normalize_newlines(text: str) -> str:
    return shared_normalize_newlines(text)


def normalize_text(text: str) -> str:
    return normalize_qiaopi_text(text)


def first_match(pattern: str, text: str, flags: int = 0) -> str:
    match = re.search(pattern, text, flags)
    if not match:
        return ""
    if match.lastindex:
        return next((group.strip() for group in match.groups() if group), "")
    return match.group(0).strip()


def unique_join(values: Iterable[str], sep: str = "；") -> str:
    seen: set[str] = set()
    result: list[str] = []
    for value in values:
        item = value.strip()
        if not item or item in seen:
            continue
        seen.add(item)
        result.append(item)
    return sep.join(result)


def parse_metadata_block(body_clean: str) -> dict[str, str]:
    metadata: dict[str, str] = {}
    for line in body_clean.split("\n")[:8]:
        for label, field in METADATA_LABELS.items():
            pattern = rf"^{re.escape(label)}\s*[:：]\s*(.+)$"
            match = re.match(pattern, line)
            if match:
                metadata[field] = match.group(1).strip()
    return metadata


def is_metadata_line(line: str) -> bool:
    return bool(METADATA_LINE_PATTERN.match(line.strip()))


def is_placeholder_note(line: str) -> bool:
    stripped = line.strip()
    unwrapped = stripped.strip("【】[]")
    return bool(
        re.search(r"(无正文录文|未见正文录文|未录正文|缺正文|待补)", unwrapped)
        or (
            re.search(r"(仅见|仅含).*(寄批人|收批人).*(批款|日期)", unwrapped)
            and re.search(r"(正文|截图|录文|未见)", unwrapped)
        )
    )


def remove_placeholder_notes(text: str) -> str:
    return normalize_text("\n".join(line for line in text.split("\n") if not is_placeholder_note(line)))


def remove_metadata_and_placeholders(body_clean: str) -> str:
    lines: list[str] = []
    in_initial_metadata = True
    for line in body_clean.split("\n"):
        stripped = line.strip()
        if in_initial_metadata and is_metadata_line(stripped):
            continue
        if not stripped:
            in_initial_metadata = False
            lines.append("")
            continue
        in_initial_metadata = False
        if is_placeholder_note(stripped):
            continue
        lines.append(stripped)
    return normalize_text("\n".join(lines))


def split_sentences(text: str) -> list[str]:
    normalized = normalize_text(text)
    if not normalized:
        return []
    rough_parts = re.split(r"(?<=[。！？!?；;])\s*|\n+", normalized)
    sentences: list[str] = []
    for part in rough_parts:
        item = part.strip()
        if not item:
            continue
        if len(item) > 120 and "，" in item:
            sentences.extend(chunk.strip() for chunk in re.split(r"(?<=，)", item) if chunk.strip())
        else:
            sentences.append(item)
    return sentences


def first_sentence_with_keywords(sentences: list[str], keywords: Iterable[str], max_items: int = 1) -> str:
    hits = [sentence for sentence in sentences if any(keyword in sentence for keyword in keywords)]
    return unique_join(hits[:max_items])


def find_terms(text: str, terms: Iterable[str]) -> list[str]:
    return [term for term in terms if term and term in text]


def is_salutation_candidate(line: str) -> bool:
    stripped = line.strip()
    if not stripped or is_metadata_line(stripped):
        return False
    if len(stripped) > 60:
        return False
    if SIGNATURE_BODY_FEATURE_PATTERN.search(stripped):
        return False
    return bool(SALUTATION_PATTERN.search(stripped))


def is_standalone_date_line(line: str) -> bool:
    stripped = re.sub(r"\s+", "", line.strip())
    if not stripped or len(stripped) > 45:
        return False
    if SIGNATURE_BODY_FEATURE_PATTERN.search(stripped):
        return False
    has_date_marker = bool(re.search(r"(民国|民國|夏历|夏曆|旧历|旧曆|农历|農曆|西历|西曆|年|月|日|号|號|初|廿|卅)", stripped))
    return has_date_marker and bool(DATE_LINE_PATTERN.match(stripped))


def is_signature_candidate(line: str) -> bool:
    stripped = line.strip()
    if not stripped or is_metadata_line(stripped):
        return False
    if len(stripped) > 50:
        return False
    if len(stripped) > 25 and SIGNATURE_BODY_FEATURE_PATTERN.search(stripped):
        return False
    return bool(SIGNATURE_FORMULA_PATTERN.search(stripped))


def is_tail_date_or_signature_line(line: str) -> bool:
    stripped = line.strip()
    if is_standalone_date_line(stripped) or is_signature_candidate(stripped):
        return True
    if len(stripped) > 60 or SIGNATURE_BODY_FEATURE_PATTERN.search(stripped):
        return False
    has_date = bool(re.search(r"(民国|民國|夏历|夏曆|旧历|旧曆|阴历|陰曆|农历|農曆|西历|西曆|年|月|日|号|號|初|廿|卅)", stripped))
    has_signature_tail = bool(SIGNATURE_FORMULA_PATTERN.search(stripped))
    has_date_triplet = bool(re.match(r"^[零〇一壹乙二贰貳弍式两兩三叁參四肆五伍六陆陸七柒八捌九玖十拾百佰千仟万萬廿卅\d]+[，,、]", stripped))
    has_short_sender_tail = bool(
        re.search(r"[\u4e00-\u9fff]{1,8}(寄|缄|禀|上|手)$", stripped)
        or re.search(r"(缄|禀|上|手)\s*[\u4e00-\u9fff]{1,8}$", stripped)
    )
    return (has_date and (has_signature_tail or has_short_sender_tail)) or (has_date_triplet and has_short_sender_tail)


def chinese_integer_to_number(text: str) -> int | None:
    clean = re.sub(r"[^零〇一壹乙二贰貳弍式两兩三叁參四肆五伍六陆陸七柒八捌九玖十拾百佰千仟万萬廿卅\d]", "", text)
    if not clean:
        return None
    if clean.isdigit():
        return int(clean)
    if clean.startswith("廿"):
        tail = chinese_integer_to_number(clean[1:])
        return 20 + (tail or 0)
    if clean.startswith("卅"):
        tail = chinese_integer_to_number(clean[1:])
        return 30 + (tail or 0)

    total = 0
    section = 0
    number = 0
    seen = False
    for char in clean:
        if char.isdigit():
            number = int(char)
            seen = True
        elif char in CHINESE_DIGITS:
            number = CHINESE_DIGITS[char]
            seen = True
        elif char in SMALL_UNITS:
            unit = SMALL_UNITS[char]
            section += (number if number else 1) * unit
            number = 0
            seen = True
        elif char in LARGE_UNITS:
            unit = LARGE_UNITS[char]
            total += (section + number) * unit
            section = 0
            number = 0
            seen = True
    if not seen:
        return None
    return total + section + number


def normalize_currency(currency_text: str, unit_text: str) -> str:
    clean = currency_text.strip().replace("（", "(").replace("）", ")")
    if clean in MONEY_CURRENCY_ALIASES:
        return MONEY_CURRENCY_ALIASES[clean]
    restored = clean.replace("(", "（").replace(")", "）")
    if restored in MONEY_CURRENCY_ALIASES:
        return MONEY_CURRENCY_ALIASES[restored]
    if "港币" in clean:
        return "港币"
    if "国币" in clean:
        return "国币"
    if "中央" in clean or "法币" in clean or "国洋" in clean:
        return "国币"
    if "金圆" in clean or "金员" in clean:
        return "金圆券"
    if "荷银" in clean:
        return "荷银"
    if "英洋" in clean:
        return "英洋"
    if "大洋" in clean:
        return "大洋"
    if "银" in clean or "銀" in clean or "良" in clean:
        return "银元"
    return unit_text if unit_text else ""


def extract_remittance(body_clean: str, metadata: dict[str, str], sentences: list[str]) -> dict[str, object]:
    metadata_raw = metadata.get("remittance_raw", "")
    remittance_sentences = [
        sentence
        for sentence in sentences
        if re.search(r"寄上|汇上|批款|查收|收讫|付去|寄去|奉.*(?:元|银|洋)|兹.*(?:元|银|洋)", sentence)
    ]
    candidate_texts: list[tuple[int, int, str]] = []
    if metadata_raw:
        candidate_texts.append((100, 0, metadata_raw))
    scored_sentences: list[tuple[int, int, str]] = []
    for idx, sentence in enumerate(remittance_sentences):
        score = 10
        if re.search(r"兹|寄上|汇上|寄去|付去|奉|批款", sentence):
            score += 30
        if re.search(r"查收|收讫|收用|分抹", sentence):
            score += 20
        if re.search(r"大洋|荷银|英洋|港币|国币|中央币|中央法币|法币|银元|大银|大龙银|良[（(]?银|银", sentence):
            score += 10
        if re.search(r"每月|逐月|最多|只能|如旧", sentence):
            score -= 15
        candidate_texts.append((score, idx + 1, sentence))
        scored_sentences.append((score, idx, sentence))
    candidate_texts.append((0, len(candidate_texts), body_clean))
    statement = unique_join(
        sentence for _, _, sentence in sorted(scored_sentences, key=lambda item: (-item[0], item[1]))[:2]
    )

    best_match: re.Match[str] | None = None
    best_score = -1
    best_order = -1
    for score, order, candidate_text in candidate_texts:
        for candidate_match in MONEY_MENTION_PATTERN.finditer(candidate_text):
            currency_text = candidate_match.group("currency") or ""
            adjusted_score = score + (5 if currency_text else 0)
            if adjusted_score > best_score or (adjusted_score == best_score and order > best_order):
                best_match = candidate_match
                best_score = adjusted_score
                best_order = order

    amount_text = ""
    amount_number: int | float | str = ""
    currency = ""
    remittance_raw = metadata_raw
    if best_match:
        currency_text = best_match.group("currency") or ""
        unit_text = best_match.group("unit") or ""
        amount_body = best_match.group("amount") or ""
        amount_text = f"{amount_body}{unit_text}"
        parsed = chinese_integer_to_number(amount_body)
        if parsed is not None and unit_text in {"万元", "萬元"}:
            parsed *= 10000
        amount_number = parsed if parsed is not None else ""
        currency = normalize_currency(currency_text, unit_text)
        remittance_raw = metadata_raw or best_match.group(0).strip()

    receipt_instruction = first_sentence_with_keywords(
        sentences,
        ["查收", "收讫", "照收", "收明", "收用", "分抹", "交", "笑纳"],
        max_items=2,
    )
    purpose_map = {
        "家用": ["家中之用", "家用", "家需"],
        "茶果": ["茶果"],
        "升学": ["升学", "学费", "读书"],
        "赎厝": ["赎厝", "续（赎）回", "续赎", "赎回"],
        "债务": ["借", "债", "欠", "还"],
        "医药": ["病", "药", "疾"],
        "分抹": ["分抹", "交"],
    }
    purposes = [label for label, keywords in purpose_map.items() if any(keyword in statement for keyword in keywords)]
    has_remittance = int(bool(metadata_raw or statement or best_match))
    return {
        "remittance_raw": remittance_raw,
        "remittance_amount_text": amount_text,
        "remittance_amount_number": amount_number,
        "currency": currency,
        "remittance_statement": statement,
        "remittance_purpose": unique_join(purposes),
        "receipt_instruction": receipt_instruction,
        "has_remittance": has_remittance,
    }


def extract_amount_mentions(
    record_id: str,
    body_clean: str,
    primary_remittance: dict[str, object],
) -> list[dict[str, object]]:
    sentences = split_sentences(body_clean)
    mentions: list[dict[str, object]] = []
    seen: set[tuple[str, str, int]] = set()
    for sentence_index, sentence in enumerate(sentences):
        for match in MONEY_MENTION_PATTERN.finditer(sentence):
            raw_text = match.group(0).strip()
            if not raw_text:
                continue
            amount_body = match.group("amount") or ""
            unit_text = match.group("unit") or ""
            amount_text = f"{amount_body}{unit_text}"
            currency = normalize_currency(match.group("currency") or "", unit_text)
            amount_number = chinese_integer_to_number(amount_body)
            if amount_number is not None and unit_text in {"万元", "萬元"}:
                amount_number *= 10000
            key = (raw_text, sentence, sentence_index)
            if key in seen:
                continue
            seen.add(key)
            mentions.append(
                {
                    "record_id": record_id,
                    "raw_text": raw_text,
                    "amount_text": amount_text,
                    "amount_number": amount_number if amount_number is not None else "",
                    "currency": currency,
                    "sentence": sentence,
                    "is_primary_candidate": False,
                    "source_field": "body_clean",
                }
            )

    primary_raw = str(primary_remittance.get("remittance_raw", "") or "")
    primary_amount = str(primary_remittance.get("remittance_amount_text", "") or "")
    primary_currency = str(primary_remittance.get("currency", "") or "")
    best_index = -1
    best_score = 0
    for index, mention in enumerate(mentions):
        score = 0
        raw_text = str(mention["raw_text"])
        sentence = str(mention["sentence"])
        if primary_raw and raw_text in primary_raw:
            score += 100
        if primary_amount and mention["amount_text"] == primary_amount:
            score += 40
        if primary_currency and mention["currency"] == primary_currency:
            score += 20
        if re.search(r"寄上|汇上|批款|查收|收讫|付去|寄去|奉|兹", sentence):
            score += 10
        if score > best_score:
            best_score = score
            best_index = index
    if best_index >= 0:
        mentions[best_index]["is_primary_candidate"] = True
    return mentions


def find_place_aliases(text: str, include_prefix_aliases: bool = False) -> list[str]:
    places = [normalized for alias, normalized in PLACE_ALIASES.items() if alias in text]
    if include_prefix_aliases:
        for alias, normalized in PLACE_PREFIX_ALIASES.items():
            if alias == "香":
                if re.search(r"(^|[，,\s])香(\s*\[香港\]|港|回归|抵港|下汕)", text):
                    places.append(normalized)
            elif alias == "越":
                if re.search(r"(^|[，,\s「\[]|寄批人[:：]?|号|號|呈|进|進|退)越[\u4e00-\u9fff\[]", text):
                    places.append(normalized)
            elif alias in text:
                places.append(normalized)
    return places


def extract_places(text: str, sender: str, recipient: str, title: str) -> dict[str, object]:
    combined = unique_join([sender, recipient, title, text], sep="\n")
    mentions: list[str] = []
    for alias, normalized in PLACE_ALIASES.items():
        if alias in combined:
            mentions.append(f"{alias}->{normalized}" if alias != normalized else normalized)
    if re.search(r"(^|[在往赴下抵来回])叻", combined):
        mentions.append("叻->新加坡")
    if re.search(r"(^|[在往赴下抵来回])暹", combined):
        mentions.append("暹->泰国")

    title_sender_side = title.split("寄")[0] if "寄" in title else title
    sender_places = find_place_aliases(unique_join([sender, title_sender_side]), include_prefix_aliases=True)
    recipient_side = title.split("寄", 1)[1] if "寄" in title else recipient
    recipient_places = find_place_aliases(unique_join([recipient, recipient_side]), include_prefix_aliases=False)
    mentioned_places = [normalized for alias, normalized in PLACE_ALIASES.items() if alias in combined]
    if re.search(r"(^|[在往赴下抵来回])叻", combined):
        mentioned_places.append("新加坡")
    if re.search(r"(^|[在往赴下抵来回])暹", combined):
        mentioned_places.append("泰国")

    overseas_place = next((place for place in sender_places + recipient_places if place in OVERSEAS_NORMALIZED), "")
    hometown_place = next((place for place in recipient_places + sender_places if place not in OVERSEAS_NORMALIZED and place.startswith("广东")), "")
    origin_place = overseas_place or (sender_places[0] if sender_places else "")
    destination_place = hometown_place or (recipient_places[0] if recipient_places else "")
    countries = unique_join(
        COUNTRY_OR_REGION_BY_PLACE.get(place, "")
        for place in [origin_place, destination_place, overseas_place, hometown_place, *mentioned_places]
    )
    route_parts = [part for part in [origin_place, destination_place] if part]
    return {
        "origin_place": origin_place,
        "destination_place": destination_place,
        "overseas_place": overseas_place,
        "hometown_place": hometown_place,
        "place_mentions": unique_join(mentions),
        "route_text": " -> ".join(route_parts),
        "is_overseas_to_hometown": int(bool(overseas_place and hometown_place and origin_place == overseas_place)),
        "country_or_region": countries,
    }


def extract_title_sender_recipient(title: str) -> tuple[str, str]:
    if "寄" not in title:
        return "", ""
    before, after = title.rsplit("寄", 1)
    before = re.sub(r"^.*?[，,]\s*", "", before).strip()
    before = re.sub(r"^\S*?\s*\[[^\]]+\]\s*", "", before).strip()
    after = re.sub(r"[，,。].*$", "", after).strip()
    return before, after


def extract_title_parts(title: str) -> dict[str, str]:
    title = title.strip()
    date_text = ""
    first_segment_match = re.match(r"\s*(.+?)(?:，|,)\s*", title)
    if first_segment_match and re.search(r"年|月|日|号|號|初|廿|卅|[(（]\d{4}", first_segment_match.group(1)):
        date_text = first_segment_match.group(1).strip()
    elif "寄" in title:
        possible_date = title.split("寄", 1)[0]
        possible_date = re.sub(r"[^，,]+$", "", possible_date).strip("，, ")
        if re.search(r"年|月|日|号|號|初|廿|卅|[(（]\d{4}", possible_date):
            date_text = possible_date

    uncertain = title_has_multiple_send_segments(title) or bool(re.search(r"[；;].*寄", title))
    if uncertain or "寄" not in title:
        return {
            "title_date_text": date_text,
            "title_sender_text": "",
            "title_recipient_text": "",
        }

    before, after = title.rsplit("寄", 1)
    sender_text = before.strip()
    if date_text and sender_text.startswith(date_text):
        sender_text = sender_text[len(date_text) :].strip("，, ")
    else:
        sender_text = re.sub(r"^.*?[，,]\s*", "", sender_text).strip()
    recipient_text = re.sub(r"[，,。；;].*$", "", after).strip()
    return {
        "title_date_text": date_text,
        "title_sender_text": sender_text,
        "title_recipient_text": recipient_text,
    }


def clean_person_name(value: str, role: str) -> str:
    original = value.strip()
    if not original:
        return ""
    cleaned = re.sub(r"[［\[][^\]］]*[\]］]", "", original)
    cleaned = re.sub(r"[（(][^)）]*[)）]", "", cleaned)
    cleaned = cleaned.strip(" ，,。:：")
    leading_place_terms = sorted(
        set(PLACE_ALIASES) | set(PLACE_ALIASES.values()) | set(PLACE_PREFIX_ALIASES) | {"外洋", "海外"},
        key=len,
        reverse=True,
    )
    changed = True
    while changed:
        changed = False
        for term in leading_place_terms:
            if term and cleaned.startswith(term) and len(cleaned) > len(term) + 1:
                cleaned = cleaned[len(term) :].strip(" ，,。")
                changed = True
                break
    if role == "recipient":
        cleaned = re.sub(r"^.*?家(?=(慈亲|严亲|双亲|母亲|父亲|祖|岳|兄|弟|妻|大嫂|媳|吾儿|侄儿))", "", cleaned)
        cleaned = re.sub(r"^(潮汕|潮安|澄海|饶邑|汀邑|海邑|广东|汕头|南桂|大寨|西洋乡|程洋岗|程洋冈)", "", cleaned)
    cleaned = re.sub(r"(尊前|膝下|台鉴|英鉴|钧鉴|清鉴|尊鉴|赐鉴|收知|如晤|侍右|尊右)$", "", cleaned)
    cleaned = re.sub(r"(大人|少爷|先生|吾儿|贤妻|内助|收阅|收次|收鉴|如见|大鉴)$", "", cleaned)
    cleaned = cleaned.strip(" ，,。:：")
    if not cleaned or len(cleaned) < 1:
        return original
    if len(cleaned) > 12 and len(original) <= 12:
        return original
    return cleaned


KINSHIP_PRIORITY = [
    "双亲",
    "父亲",
    "母亲",
    "慈亲",
    "严亲",
    "祖父",
    "祖母",
    "外祖父母",
    "外祖父",
    "外祖母",
    "祖慈",
    "家祖",
    "岳父",
    "岳母",
    "妻",
    "贤妻",
    "内妻",
    "内助",
    "兄",
    "弟",
    "姊",
    "姐",
    "妹",
    "吾儿",
    "儿",
    "男",
    "孙",
    "外孙",
    "媳",
    "媳妇",
    "侄",
    "大嫂",
    "嫂",
    "叔",
    "伯",
    "婶",
    "姑",
    "妗",
]


def select_primary_kinship(kinship_terms: str, recipient: str, opening: str) -> str:
    source = unique_join([kinship_terms, recipient, opening])
    for term in KINSHIP_PRIORITY:
        if term in source:
            return term
    return ""


def relation_confidence_label(
    relationship_type: str,
    recipient_role: str,
    kinship_terms: str,
    opening_salutation: str,
    self_reference: str,
    title_sender_text: str,
    title_recipient_text: str,
) -> str:
    if not relationship_type or relationship_type == "unknown":
        return ""
    strong_signals = sum(bool(value) for value in [recipient_role, kinship_terms, opening_salutation, self_reference])
    if strong_signals >= 3:
        return "high"
    if strong_signals >= 2:
        return "medium"
    if title_sender_text or title_recipient_text or strong_signals == 1:
        return "low"
    return ""


def normalized_place_items_from_row(row: dict[str, object]) -> list[dict[str, str]]:
    fields = {
        "title_reference": str(row.get("title_reference", "") or ""),
        "sender": str(row.get("sender", "") or ""),
        "recipient": str(row.get("recipient", "") or ""),
        "body_clean": str(row.get("body_clean", "") or ""),
        "place_mentions": str(row.get("place_mentions", "") or ""),
    }
    items: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for source_field, text in fields.items():
        if not text:
            continue
        if source_field == "place_mentions":
            for part in text.split("；"):
                if not part.strip():
                    continue
                alias_text, normalized = (part.split("->", 1) + [""])[:2] if "->" in part else (part, part)
                normalized = normalized or alias_text
                key = (alias_text, normalized, source_field)
                if key not in seen:
                    seen.add(key)
                    items.append(
                        {
                            "alias_text": alias_text,
                            "normalized_place": normalized,
                            "country_or_region": COUNTRY_OR_REGION_BY_PLACE.get(normalized, ""),
                            "source_field": source_field,
                        }
                    )
            continue
        for alias, normalized in PLACE_ALIASES.items():
            if alias in text:
                key = (alias, normalized, source_field)
                if key not in seen:
                    seen.add(key)
                    items.append(
                        {
                            "alias_text": alias,
                            "normalized_place": normalized,
                            "country_or_region": COUNTRY_OR_REGION_BY_PLACE.get(normalized, ""),
                            "source_field": source_field,
                        }
                    )
        if re.search(r"(^|[在往赴下抵来回])叻", text):
            key = ("叻", "新加坡", source_field)
            if key not in seen:
                seen.add(key)
                items.append(
                    {
                        "alias_text": "叻",
                        "normalized_place": "新加坡",
                        "country_or_region": "新加坡",
                        "source_field": source_field,
                    }
                )
    return items


def place_mentions_normalized(row: dict[str, object]) -> str:
    normalized_values = [item["normalized_place"] for item in normalized_place_items_from_row(row)]
    normalized_values.extend(
        str(row.get(field, "") or "")
        for field in ["origin_place", "destination_place", "overseas_place", "hometown_place"]
        if row.get(field)
    )
    return unique_join(normalized_values)


def count_ocr_uncertainty(text: str) -> int:
    count = text.count("□") + text.count("�")
    for marker in ["缺字", "待校", "无法辨识", "疑为", "未辨", "不清"]:
        count += text.count(marker)
    return count


def generate_retrieval_keywords(row: dict[str, object]) -> str:
    text_parts = [
        str(row.get("title_reference", "")),
        str(row.get("sender", "")),
        str(row.get("recipient", "")),
        str(row.get("sender_name_clean", "")),
        str(row.get("recipient_name_clean", "")),
        str(row.get("kinship_terms", "")),
        str(row.get("primary_kinship", "")),
        str(row.get("place_mentions_normalized", "")),
        str(row.get("remittance_raw", "")),
        str(row.get("currency", "")),
        str(row.get("theme_tags", "")),
        str(row.get("style_keywords", "")),
        str(row.get("body_core", "")),
    ]
    keywords: list[str] = []

    def add(value: str) -> None:
        value = value.strip(" ，,。:：；;、\n\t")
        if value and value not in keywords:
            keywords.append(value)

    for field in text_parts[:11]:
        for part in re.split(r"[；;\s\n]+", field):
            add(part)
    for term in [*KINSHIP_TERMS, *PLACE_ALIASES.keys(), *PLACE_ALIASES.values(), *MONEY_CURRENCY_ALIASES.keys(), *COMMON_FORMULAE]:
        if any(term in part for part in text_parts):
            add(term)
            if term in PLACE_ALIASES:
                add(PLACE_ALIASES[term])
    body_core = str(row.get("body_core", ""))
    try:
        import jieba  # type: ignore

        tokens = jieba.lcut(unique_join(text_parts, sep="\n"))
    except Exception:
        tokens = re.findall(r"[\u4e00-\u9fff]{2,6}|[A-Za-z0-9_]{2,}", unique_join(text_parts, sep="\n"))
    stopwords = {"一个", "一切", "以及", "但是", "不可", "此致", "敬请", "顺此", "查收", "大人"}
    for token in tokens:
        token = token.strip()
        if len(token) < 2 or token in stopwords:
            continue
        if re.fullmatch(r"\d+", token):
            continue
        if token in body_core or any(token in part for part in text_parts[:-1]):
            add(token)
        if len(keywords) >= 80:
            break
    return "；".join(keywords[:80])


def detect_role(text: str, self_reference: str = "") -> str:
    source = unique_join([text, self_reference])
    if any(term in source for term in ["祖父", "祖母", "外祖父", "外祖母", "外祖父母", "祖慈", "家祖", "岳祖母"]):
        return "grandparent"
    if any(term in source for term in ["岳父", "岳母", "岳慈亲"]):
        return "parent_in_law"
    if any(term in source for term in ["父亲", "母亲", "双亲", "慈亲", "严亲", "阿嬷", "姆母", "母兴", "花灯"]):
        return "parent"
    if any(term in source for term in ["妻", "贤妻", "内助", "内妻", "荆妻", "爱卿"]):
        return "spouse"
    if any(term in source for term in ["嫂", "大嫂"]):
        return "sibling_in_law"
    if any(term in source for term in ["兄", "弟", "姊", "姐", "妹"]):
        return "sibling"
    if any(term in source for term in ["叔", "伯", "婶", "姑", "妗", "婶母", "叔父", "伯母"]):
        return "elder_relative"
    if "侄" in source:
        return "nephew_or_niece"
    if "媳" in source:
        return "daughter_in_law"
    if any(term in source for term in ["吾儿", "儿知悉", "侄儿"]):
        return "child"
    if any(term == self_reference for term in ["儿", "男", "小儿"]):
        return "child"
    if any(term == self_reference for term in ["孙", "外孙"]):
        return "grandchild"
    if self_reference == "婿":
        return "son_in_law"
    if self_reference == "夫":
        return "husband"
    if "先生" in source or "台鉴" in source:
        return "social_or_business_contact"
    return ""


def infer_relationship(recipient_role: str, sender_role: str, recipient: str, self_reference: str) -> str:
    if recipient_role == "parent" or sender_role == "child" or self_reference in {"儿", "小儿", "男"}:
        return "child_to_parent"
    if recipient_role == "grandparent" or sender_role == "grandchild" or self_reference in {"孙", "外孙"}:
        return "grandchild_to_grandparent"
    if recipient_role == "parent_in_law" or sender_role == "son_in_law" or self_reference == "婿":
        return "son_in_law_to_parent_in_law"
    if recipient_role == "spouse":
        return "spouse_to_spouse"
    if sender_role == "husband":
        return "spouse_to_spouse"
    if recipient_role == "sibling" or sender_role == "sibling":
        return "sibling_or_same_generation_kin"
    if recipient_role == "sibling_in_law":
        return "in_law_or_affinal_kin"
    if recipient_role == "elder_relative":
        return "younger_to_elder_kin"
    if "侄" in recipient or "侄" in self_reference:
        return "uncle_nephew_or_nephew_kin"
    if recipient_role == "daughter_in_law":
        return "parent_to_child_in_law"
    if recipient_role == "child":
        return "parent_to_child"
    if recipient_role == "social_or_business_contact":
        return "social_or_business"
    return "unknown"


def extract_signature(lines: list[str]) -> tuple[str, str, str]:
    non_empty = [line.strip() for line in lines if line.strip()]
    if not non_empty:
        return "", "", ""
    postscript_start = next((i for i, line in enumerate(non_empty) if re.search(r"另禀者|再禀|又及|附言", line)), None)
    searchable = non_empty[:postscript_start] if postscript_start is not None else non_empty
    tail = searchable[-8:]
    date_line = ""
    signature = ""
    for line in reversed(tail):
        if is_standalone_date_line(line) or re.search(r"(民国|夏历|年|月|日|号|號|初|廿|卅)", line):
            if is_signature_candidate(line):
                signature = signature or line
            elif is_standalone_date_line(line):
                date_line = date_line or line
    for line in reversed(tail):
        if is_signature_candidate(line):
            signature = line
            break
    if len(signature) > 25 and SIGNATURE_BODY_FEATURE_PATTERN.search(signature):
        signature = ""
    date_signature_line = unique_join([signature, date_line], sep="\n")
    closing_formula = first_match(
        r"(谨禀|谨叩|谨上|叩上|手书|手泐|拜上|禀上|敬上|上言|拜启|具禀|泐|顿|草|字|启|托|缄)$",
        signature,
    )
    return signature, date_signature_line, closing_formula


def extract_postscript(text: str) -> str:
    match = re.search(r"(另禀者|再禀|又及|附言).+$", text, flags=re.S)
    return normalize_text(match.group(0)) if match else ""


def parse_letter_structure(body_without_metadata: str, sentences: list[str]) -> dict[str, str]:
    lines = [line.strip() for line in body_without_metadata.split("\n") if line.strip()]
    opening_salutation = ""
    for line in lines[:5]:
        if is_salutation_candidate(line):
            opening_salutation = line
            break

    opening_formula = first_sentence_with_keywords(sentences[:4], ["敬禀者", "叩禀者", "启者", "禀者"], max_items=1)
    opening_greeting = first_sentence_with_keywords(sentences[:8], ["平安", "安好", "清泰", "起居", "痊安", "福安"], max_items=1)
    safety_report = first_sentence_with_keywords(sentences, SAFETY_FORMULAE, max_items=2)
    family_care_statement = first_sentence_with_keywords(
        sentences,
        ["保重", "身体", "安康", "家中", "家务", "父", "母", "妻", "儿", "弟妹", "合家"],
        max_items=2,
    )
    instruction_statement = first_sentence_with_keywords(sentences, INSTRUCTION_FORMULAE, max_items=2)
    closing_greeting = first_sentence_with_keywords(sentences[-6:], ["此候", "顺候", "并请", "此问", "此致", "顺此"], max_items=1)
    signature, date_signature_line, closing_formula = extract_signature(lines)
    postscript = extract_postscript(body_without_metadata)

    main_lines = list(lines)
    date_signature_parts = {part.strip() for part in date_signature_line.split("\n") if part.strip()}
    for removable in [opening_salutation, signature]:
        if removable in main_lines:
            main_lines.remove(removable)
    if postscript:
        postscript_first = postscript.split("\n", 1)[0]
        main_lines = [line for line in main_lines if line != postscript_first and line not in postscript.split("\n")]
    tail_start = max(0, len(main_lines) - 8)
    cleaned_main_lines: list[str] = []
    for idx, line in enumerate(main_lines):
        in_tail = idx >= tail_start
        if line in date_signature_parts or line == date_signature_line:
            continue
        if is_standalone_date_line(line):
            continue
        if in_tail and is_tail_date_or_signature_line(line):
            continue
        if line == closing_greeting and len(line) <= 8:
            continue
        if line == closing_formula:
            continue
        cleaned_main_lines.append(line)
    main_lines = cleaned_main_lines
    main_message = normalize_text("\n".join(main_lines))
    return {
        "opening_salutation": opening_salutation,
        "opening_formula": opening_formula,
        "opening_greeting": opening_greeting,
        "main_message": main_message,
        "safety_report": safety_report,
        "family_care_statement": family_care_statement,
        "instruction_statement": instruction_statement,
        "closing_greeting": closing_greeting,
        "closing_formula": closing_formula,
        "date_signature_line": date_signature_line,
        "signature": signature,
        "postscript": postscript,
    }


def extract_people(metadata: dict[str, str], title: str, body_without_metadata: str, structure: dict[str, str]) -> dict[str, str]:
    title_sender, title_recipient = extract_title_sender_recipient(title)
    sender = metadata.get("sender", "") or title_sender
    recipient = metadata.get("recipient", "") or title_recipient
    opening = structure.get("opening_salutation", "")
    if not recipient and opening:
        recipient = re.sub(r"[:：].*$", "", opening).strip()

    signature = structure.get("signature", "")
    self_reference = first_match(r"^(小儿|外孙|孙|儿|男|愚弟|愚姆|愚|弟|侄|媳|婿|夫|予|俺)", signature)
    if not self_reference:
        self_reference = first_match(r"(小儿|外孙|孙|儿|男|愚弟|愚姆|愚|弟|侄|媳|婿|夫|予|俺)(?=[\u4e00-\u9fff]{1,4}(谨|叩|启|托|手书|禀|上))", signature)
    if not sender and signature:
        sender = re.sub(r"(谨禀|谨叩|谨上|叩上|手书|启|托|拜上).*$", "", signature).strip()

    relation_text = unique_join([sender, recipient, opening, body_without_metadata, self_reference], sep="\n")
    kinship_terms = find_terms(relation_text, KINSHIP_TERMS)
    member_mentions = re.findall(
        r"[\u4e00-\u9fff]{1,4}(?:祖父|祖母|外祖父母|外祖父|外祖母|岳母|岳父|婶母|伯母|叔父|大嫂|兄|弟|妹|姊|姐|叔|伯|婶|母|父|妻|儿|孙|侄|媳|嫂|姑|妗|内侄)",
        relation_text,
    )
    recipient_honorific = unique_join(find_terms(unique_join([recipient, opening]), HONORIFIC_TERMS))
    sender_role = detect_role(sender, self_reference)
    recipient_role = detect_role(unique_join([recipient, opening]))
    relationship_type = infer_relationship(recipient_role, sender_role, recipient, self_reference)
    return {
        "sender": sender,
        "recipient": recipient,
        "sender_role": sender_role,
        "recipient_role": recipient_role,
        "relationship_type": relationship_type,
        "kinship_terms": unique_join(kinship_terms),
        "family_members_mentioned": unique_join(member_mentions[:12]),
        "recipient_honorific": recipient_honorific,
        "self_reference": self_reference,
    }


def extract_date(metadata: dict[str, str], title: str, structure: dict[str, str]) -> dict[str, object]:
    date_text = metadata.get("date_text", "")
    title_date_match = re.match(r"\s*(.+?)(?:，|,)\s*", title)
    title_date_text = title_date_match.group(1).strip() if title_date_match else ""
    if not date_text and re.search(r"年|月|日|号|號|初|廿|卅|[(（]\d{4}", title_date_text):
        date_text = title_date_text
    if not date_text:
        date_text = first_match(r"(民国[零〇一二两三四五六七八九十廿卅\d]+年[^,，。]*)", title)
    if not date_text:
        date_text = first_match(r"([甲乙丙丁戊己庚辛壬癸][^,，。]*?(?:日|号|號))", title)
    written_date_text = structure.get("date_signature_line", "")

    combined = unique_join([date_text, written_date_text, title], sep=" ")
    normalized_combined = combined.replace("民國", "民国").replace("冊", "卅")
    approximate_year = bool(re.search(r"\d{4}\s*(?:后|後|前|左右|约|約)", normalized_combined))
    year_raw = (
        first_match(r"(民国[零〇一二两三四五六七八九十廿卅\d]+年)", normalized_combined)
        or first_match(r"[(（]\s*(\d{4}\s*(?:后|後|前|左右|约|約)?)\s*[)）]", normalized_combined)
        or first_match(r"(\d{4})年", normalized_combined)
        or first_match(rf"([{STEMS_BRANCHES}]年)", normalized_combined)
    )
    year_normalized: int | str = ""
    minguo_match = re.search(r"民国([零〇一二两三四五六七八九十廿卅\d]+)年", normalized_combined)
    if minguo_match:
        year_number = chinese_integer_to_number(minguo_match.group(1))
        if year_number:
            year_normalized = 1911 + year_number
    if year_normalized == "":
        western_match = re.search(r"[(（]\s*(\d{4})\s*[)）]", normalized_combined)
        if not western_match:
            western_match = re.search(r"(\d{4})年", normalized_combined)
        if western_match:
            year = int(western_match.group(1))
            if 1800 <= year <= 2100 and not approximate_year:
                year_normalized = year

    month_text = first_match(r"([正一二两三四五六七八九十冬腊十一十二]+月)", normalized_combined)
    day_text = first_match(r"((?:初|廿|卅)?[一二三四五六七八九十两\d]{1,3}(?:日|号|號))", normalized_combined)

    if re.search(r"夏历|农历|旧历", normalized_combined) and re.search(r"民国|\d{4}", normalized_combined):
        lunar_or_solar = "mixed"
    elif re.search(r"夏历|农历|旧历|初|阳月", normalized_combined):
        lunar_or_solar = "lunar"
    elif re.search(r"民国|\d{4}|号|號", normalized_combined):
        lunar_or_solar = "solar"
    else:
        lunar_or_solar = ""

    if "民国" in normalized_combined:
        era_text = "民国"
    elif re.search(rf"[{STEMS_BRANCHES}]年|^[{STEMS_BRANCHES}]", normalized_combined):
        era_text = "干支或传统纪年"
    elif year_normalized:
        era_text = "西历"
    else:
        era_text = ""

    if year_normalized:
        confidence = "high"
    elif year_raw or (month_text and day_text):
        confidence = "medium"
    elif date_text:
        confidence = "low"
    else:
        confidence = ""

    return {
        "date_text": date_text,
        "written_date_text": written_date_text,
        "year_raw": year_raw,
        "year_normalized": year_normalized,
        "month_text": month_text,
        "day_text": day_text,
        "lunar_or_solar": lunar_or_solar,
        "era_text": era_text,
        "date_confidence": confidence,
    }


def classify_themes(text: str, has_remittance: int, body_core: str = "") -> dict[str, object]:
    fields: dict[str, object] = {}
    active_tags: list[str] = []
    for field, keywords in THEME_KEYWORDS.items():
        value = int(any(keyword in text for keyword in keywords))
        if field == "theme_remittance" and has_remittance:
            value = 1
        fields[field] = value
        if value:
            active_tags.append(field.replace("theme_", ""))
    main_intent = ""
    for field, intent in INTENT_PRIORITY:
        if field != "theme_remittance" and any(keyword in body_core for keyword in MAIN_INTENT_BODY_KEYWORDS.get(field, [])):
            main_intent = intent
            break
    if not main_intent and has_remittance:
        main_intent = "remittance"
    if not main_intent:
        for field, intent in INTENT_PRIORITY:
            if field != "theme_remittance" and fields.get(field):
                main_intent = intent
                break
    fields["main_intent"] = main_intent
    fields["theme_tags"] = ";".join(active_tags)
    return fields


def extract_style(text: str, structure: dict[str, str], remittance_statement: str, safety_report: str, instruction_statement: str) -> dict[str, object]:
    qiaopi_opening = unique_join(find_terms(text, OPENING_PATTERNS))
    qiaopi_closing = unique_join(find_terms(text, CLOSING_PATTERNS))
    honorifics = unique_join(find_terms(text, HONORIFIC_TERMS))
    humble = unique_join(find_terms(text, HUMBLE_PATTERNS))
    classical = unique_join(find_terms(text, CLASSICAL_WORDS))
    common = unique_join(find_terms(text, COMMON_FORMULAE))
    remittance_formulae = unique_join(find_terms(text, REMITTANCE_FORMULAE))
    safety_formulae = unique_join(find_terms(text, SAFETY_FORMULAE))
    instruction_formulae = unique_join(find_terms(text, INSTRUCTION_FORMULAE))
    style_keywords = unique_join(
        [
            *qiaopi_opening.split("；"),
            *qiaopi_closing.split("；"),
            *honorifics.split("；"),
            *humble.split("；"),
            *classical.split("；"),
            *common.split("；"),
        ]
    )
    style_count = len([item for item in style_keywords.split("；") if item])
    if style_count >= 10:
        style_strength = 3
    elif style_count >= 5:
        style_strength = 2
    elif style_count >= 1:
        style_strength = 1
    else:
        style_strength = 0

    style_reference_text = unique_join(
        [
            structure.get("opening_salutation", ""),
            structure.get("opening_formula", ""),
            safety_report,
            remittance_statement,
            instruction_statement,
            structure.get("closing_greeting", ""),
            structure.get("signature", ""),
        ],
        sep="\n",
    )
    return {
        "qiaopi_opening_pattern": qiaopi_opening,
        "qiaopi_closing_pattern": qiaopi_closing,
        "honorific_phrases": honorifics,
        "humble_phrases": humble,
        "classical_words": classical,
        "common_formulae": common,
        "remittance_formulae": remittance_formulae,
        "safety_formulae": safety_formulae,
        "instruction_formulae": instruction_formulae,
        "style_keywords": style_keywords,
        "style_strength": style_strength,
        "style_reference_text": style_reference_text,
    }


def build_evidence_and_rag(row: dict[str, object]) -> dict[str, str]:
    entity_parts = [
        f"sender={row.get('sender', '')}",
        f"recipient={row.get('recipient', '')}",
        f"relationship={row.get('relationship_type', '')}",
        f"date={row.get('date_text', '')}",
        f"money={row.get('remittance_amount_text', '')}",
        f"currency={row.get('currency', '')}",
        f"places={row.get('place_mentions', '')}",
        f"kinship={row.get('kinship_terms', '')}",
    ]
    evidence_entities = "；".join(part for part in entity_parts if not part.endswith("="))
    retrieval_text = unique_join(
        [
            str(row.get("title_reference", "")),
            str(row.get("sender", "")),
            str(row.get("recipient", "")),
            str(row.get("date_text", "")),
            str(row.get("remittance_raw", "")),
            str(row.get("place_mentions", "")),
            str(row.get("theme_tags", "")),
            str(row.get("body_core", "")),
        ],
        sep="\n",
    )
    core = str(row.get("body_core", ""))
    summary_core = re.sub(r"\s+", "", core)[:120]
    rag_summary_text = unique_join(
        [
            f"题名：{row.get('title_reference', '')}",
            f"寄批人：{row.get('sender', '')}" if row.get("sender") else "",
            f"收批人：{row.get('recipient', '')}" if row.get("recipient") else "",
            f"日期：{row.get('date_text', '')}" if row.get("date_text") else "",
            f"汇款：{row.get('remittance_raw', '')}" if row.get("remittance_raw") else "",
            f"主题：{row.get('theme_tags', '')}" if row.get("theme_tags") else "",
            f"正文摘录：{summary_core}" if summary_core else "",
        ]
    )
    return {
        "evidence_opening": str(row.get("opening_salutation", "")) or str(row.get("opening_formula", "")),
        "evidence_safety": str(row.get("safety_report", "")),
        "evidence_remittance": str(row.get("remittance_statement", "")) or str(row.get("remittance_raw", "")),
        "evidence_family_care": str(row.get("family_care_statement", "")),
        "evidence_instruction": str(row.get("instruction_statement", "")),
        "evidence_closing": unique_join([str(row.get("closing_greeting", "")), str(row.get("signature", ""))]),
        "evidence_entities": evidence_entities,
        "retrieval_text": retrieval_text,
        "rag_summary_text": rag_summary_text,
    }


def assess_quality(body_without_metadata: str, body_clean: str, quality_flags: str, review_status: str) -> dict[str, object]:
    placeholder_only = not body_without_metadata or bool(re.search(r"仅见|未见正文|未录正文|缺正文|待补", body_clean))
    has_full_text = int(not placeholder_only and len(re.sub(r"\s+", "", body_without_metadata)) >= 30)
    uncertain = bool(re.search(r"□|缺字|待校|无法辨识", body_clean))
    if not has_full_text:
        text_quality_level = "metadata_only"
    elif quality_flags or uncertain:
        text_quality_level = "needs_review"
    elif len(body_without_metadata) >= 180:
        text_quality_level = "high"
    elif len(body_without_metadata) >= 80:
        text_quality_level = "medium"
    else:
        text_quality_level = "low"
    review_needed = int(
        bool(quality_flags)
        or uncertain
        or not has_full_text
    )
    return {
        "has_full_text": has_full_text,
        "text_quality_level": text_quality_level,
        "review_needed": review_needed,
    }


def title_has_multiple_send_segments(title: str) -> bool:
    return title.count("寄") > 1 or bool(re.search(r"寄.+寄|；.*寄|;.*寄", title))


def update_extraction_review_needed(
    quality: dict[str, object],
    title: str,
    people: dict[str, str],
    dates: dict[str, object],
    remittance: dict[str, object],
) -> None:
    key_field_missing = not people.get("sender") or not people.get("recipient") or not dates.get("date_text")
    remittance_missing_amount = bool(remittance.get("has_remittance")) and not remittance.get("remittance_amount_text")
    extraction_review_needed = bool(quality.get("review_needed")) or title_has_multiple_send_segments(title) or key_field_missing or remittance_missing_amount
    quality["review_needed"] = int(extraction_review_needed)


def build_wide_row(source_row: pd.Series, row_number: int) -> dict[str, object]:
    body_raw = normalize_newlines(safe_text(source_row.get("source_text_ocr", ""))).strip()
    body_clean = remove_placeholder_notes(normalize_text(body_raw))
    metadata = parse_metadata_block(body_clean)
    body_without_metadata = remove_metadata_and_placeholders(body_clean)
    sentences = split_sentences(body_without_metadata)
    structure = parse_letter_structure(body_without_metadata, sentences)
    body_core = structure["main_message"] or body_without_metadata

    title = safe_text(source_row.get("title_reference", "")).strip()
    title_parts = extract_title_parts(title)
    source_index_value = source_row.get("item_no", row_number + 1)
    source_index = int(source_index_value) if not pd.isna(source_index_value) else row_number + 1
    record_id = safe_text(source_row.get("record_id", "")).strip() or f"CSQP-SFHC-TEXT-{source_index:03d}"
    page_reference = unique_join(
        [
            f"pdf_pages={safe_text(source_row.get('pdf_pages', '')).strip()}",
            f"book_pages={safe_text(source_row.get('book_pages', '')).strip()}",
        ]
    )
    quality = assess_quality(
        body_without_metadata,
        body_clean,
        safe_text(source_row.get("quality_flags", "")).strip(),
        safe_text(source_row.get("review_status", "")).strip(),
    )
    remittance = extract_remittance(body_clean, metadata, sentences)
    amount_mentions = extract_amount_mentions(record_id, body_clean, remittance)
    people = extract_people(metadata, title, body_without_metadata, structure)
    places = extract_places(body_without_metadata, people["sender"], people["recipient"], title)
    dates = extract_date(metadata, title, structure)
    update_extraction_review_needed(quality, title, people, dates, remittance)
    themes = classify_themes(unique_join([title, body_clean], sep="\n"), int(remittance["has_remittance"]), body_core)
    style = extract_style(
        body_without_metadata,
        structure,
        str(remittance["remittance_statement"]),
        str(structure["safety_report"]),
        str(structure["instruction_statement"]),
    )

    row: dict[str, object] = {
        "record_id": record_id,
        "source_index": source_index,
        "title_reference": title,
        "page_reference": page_reference,
        **quality,
        "body_raw": body_raw,
        "body_clean": body_clean,
        "body_without_metadata": body_without_metadata,
        "body_core": body_core,
        "body_length": len(body_clean),
        "line_count": len([line for line in body_clean.split("\n") if line.strip()]),
        "sentence_count": len(sentences),
        **people,
        **places,
        **dates,
        **remittance,
        **structure,
        **themes,
        **style,
        **title_parts,
        "body_core_length": len(re.sub(r"\s+", "", body_core)),
        "ocr_uncertain_count": count_ocr_uncertainty(body_clean),
        "sender_name_clean": clean_person_name(people["sender"], "sender"),
        "recipient_name_clean": clean_person_name(people["recipient"], "recipient"),
        "primary_kinship": select_primary_kinship(
            people["kinship_terms"],
            people["recipient"],
            structure["opening_salutation"],
        ),
        "relation_confidence": relation_confidence_label(
            people["relationship_type"],
            people["recipient_role"],
            people["kinship_terms"],
            structure["opening_salutation"],
            people["self_reference"],
            title_parts["title_sender_text"],
            title_parts["title_recipient_text"],
        ),
        "remittance_mentions_json": json.dumps(amount_mentions, ensure_ascii=False),
    }
    row["place_mentions_normalized"] = place_mentions_normalized(row)
    row.update(build_evidence_and_rag(row))
    row["retrieval_keywords"] = generate_retrieval_keywords(row)
    return {column: row.get(column, "") for column in WIDE_COLUMNS}


def build_wide_table(frame: pd.DataFrame) -> pd.DataFrame:
    frame = frame.copy()
    frame.columns = [str(column).strip() for column in frame.columns]
    rows = [build_wide_row(source_row, idx) for idx, source_row in frame.iterrows()]
    wide = pd.DataFrame(rows, columns=WIDE_COLUMNS)
    return wide.fillna("")


AMOUNT_MENTION_COLUMNS = [
    "record_id",
    "mention_id",
    "raw_text",
    "amount_text",
    "amount_number",
    "currency",
    "sentence",
    "is_primary_candidate",
    "source_field",
]

PLACE_MENTION_COLUMNS = [
    "record_id",
    "mention_id",
    "alias_text",
    "normalized_place",
    "country_or_region",
    "source_field",
]

ENTITY_MENTION_COLUMNS = [
    "record_id",
    "mention_id",
    "entity_type",
    "entity_text",
    "normalized_text",
    "source_field",
    "confidence",
]

EVIDENCE_SPAN_COLUMNS = [
    "record_id",
    "evidence_id",
    "evidence_type",
    "evidence_text",
    "source_column",
    "start_char",
    "end_char",
]


def parse_remittance_mentions_json(value: object) -> list[dict[str, object]]:
    if not value:
        return []
    try:
        parsed = json.loads(str(value))
    except json.JSONDecodeError:
        return []
    return parsed if isinstance(parsed, list) else []


def build_amount_mentions_table(wide: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, source_row in wide.iterrows():
        record_id = str(source_row["record_id"])
        mentions = parse_remittance_mentions_json(source_row.get("remittance_mentions_json", ""))
        for index, mention in enumerate(mentions, start=1):
            rows.append(
                {
                    "record_id": record_id,
                    "mention_id": f"{record_id}-AMT-{index:03d}",
                    "raw_text": mention.get("raw_text", ""),
                    "amount_text": mention.get("amount_text", ""),
                    "amount_number": mention.get("amount_number", ""),
                    "currency": mention.get("currency", ""),
                    "sentence": mention.get("sentence", ""),
                    "is_primary_candidate": int(bool(mention.get("is_primary_candidate"))),
                    "source_field": mention.get("source_field", "body_clean"),
                }
            )
    return pd.DataFrame(rows, columns=AMOUNT_MENTION_COLUMNS)


def build_place_mentions_table(wide: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, source_row in wide.iterrows():
        record = source_row.fillna("").to_dict()
        record_id = str(record["record_id"])
        items = normalized_place_items_from_row(record)
        for index, item in enumerate(items, start=1):
            rows.append(
                {
                    "record_id": record_id,
                    "mention_id": f"{record_id}-PLC-{index:03d}",
                    "alias_text": item.get("alias_text", ""),
                    "normalized_place": item.get("normalized_place", ""),
                    "country_or_region": item.get("country_or_region", ""),
                    "source_field": item.get("source_field", ""),
                }
            )
    return pd.DataFrame(rows, columns=PLACE_MENTION_COLUMNS)


def add_entity_row(
    rows: list[dict[str, object]],
    record_id: str,
    entity_type: str,
    entity_text: str,
    normalized_text: str,
    source_field: str,
    confidence: float,
) -> None:
    entity_text = str(entity_text or "").strip()
    normalized_text = str(normalized_text or "").strip()
    if not entity_text:
        return
    rows.append(
        {
            "record_id": record_id,
            "mention_id": f"{record_id}-ENT-{len(rows) + 1:04d}",
            "entity_type": entity_type,
            "entity_text": entity_text,
            "normalized_text": normalized_text or entity_text,
            "source_field": source_field,
            "confidence": confidence,
        }
    )


def build_entity_mentions_table(wide: pd.DataFrame, place_mentions: pd.DataFrame, amount_mentions: pd.DataFrame) -> pd.DataFrame:
    rows: list[dict[str, object]] = []
    for _, source_row in wide.iterrows():
        record = source_row.fillna("").to_dict()
        record_id = str(record["record_id"])
        add_entity_row(rows, record_id, "person", record.get("sender", ""), record.get("sender_name_clean", ""), "sender", 0.85)
        add_entity_row(rows, record_id, "person", record.get("recipient", ""), record.get("recipient_name_clean", ""), "recipient", 0.85)
        for term in str(record.get("kinship_terms", "")).split("；"):
            add_entity_row(rows, record_id, "kinship", term, term, "kinship_terms", 0.8)
        for field in ["date_text", "written_date_text", "year_raw"]:
            add_entity_row(rows, record_id, "date", record.get(field, ""), record.get(field, ""), field, 0.75)
        for field in ["common_formulae", "remittance_formulae", "safety_formulae", "instruction_formulae"]:
            for formula in str(record.get(field, "")).split("；"):
                add_entity_row(rows, record_id, "style_formula", formula, formula, field, 0.7)

    for _, place in place_mentions.iterrows():
        add_entity_row(
            rows,
            str(place["record_id"]),
            "place",
            place.get("alias_text", ""),
            place.get("normalized_place", ""),
            str(place.get("source_field", "place_mentions")),
            0.8,
        )
    for _, money in amount_mentions.iterrows():
        add_entity_row(
            rows,
            str(money["record_id"]),
            "money",
            money.get("raw_text", ""),
            money.get("amount_text", ""),
            str(money.get("source_field", "body_clean")),
            0.9 if int(money.get("is_primary_candidate", 0)) else 0.7,
        )
    return pd.DataFrame(rows, columns=ENTITY_MENTION_COLUMNS)


def locate_evidence_span(body_clean: str, evidence_text: str) -> tuple[str, str, str]:
    if not evidence_text:
        return "", "", ""
    start = body_clean.find(evidence_text)
    if start >= 0:
        return "body_clean", str(start), str(start + len(evidence_text))
    for part in re.split(r"[；\n]+", evidence_text):
        part = part.strip()
        if not part:
            continue
        start = body_clean.find(part)
        if start >= 0:
            return "body_clean", str(start), str(start + len(part))
    return "", "", ""


def build_evidence_spans_table(wide: pd.DataFrame) -> pd.DataFrame:
    evidence_columns = [
        "evidence_opening",
        "evidence_safety",
        "evidence_remittance",
        "evidence_family_care",
        "evidence_instruction",
        "evidence_closing",
        "evidence_entities",
    ]
    rows: list[dict[str, object]] = []
    for _, source_row in wide.iterrows():
        record = source_row.fillna("").to_dict()
        record_id = str(record["record_id"])
        body_clean = str(record.get("body_clean", ""))
        for column in evidence_columns:
            evidence_text = str(record.get(column, "") or "").strip()
            if not evidence_text:
                continue
            source_column, start_char, end_char = locate_evidence_span(body_clean, evidence_text)
            rows.append(
                {
                    "record_id": record_id,
                    "evidence_id": f"{record_id}-EVID-{len(rows) + 1:04d}",
                    "evidence_type": column.replace("evidence_", ""),
                    "evidence_text": evidence_text,
                    "source_column": source_column or column,
                    "start_char": start_char,
                    "end_char": end_char,
                }
            )
    return pd.DataFrame(rows, columns=EVIDENCE_SPAN_COLUMNS)


def build_auxiliary_tables(wide: pd.DataFrame) -> dict[str, pd.DataFrame]:
    amount_mentions = build_amount_mentions_table(wide)
    place_mentions = build_place_mentions_table(wide)
    entity_mentions = build_entity_mentions_table(wide, place_mentions, amount_mentions)
    evidence_spans = build_evidence_spans_table(wide)
    return {
        "qiaopi_amount_mentions.csv": amount_mentions,
        "qiaopi_place_mentions.csv": place_mentions,
        "qiaopi_entity_mentions.csv": entity_mentions,
        "qiaopi_evidence_spans.csv": evidence_spans,
    }


def write_auxiliary_outputs(auxiliary_tables: dict[str, pd.DataFrame], output_dir: Path) -> dict[str, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}
    for filename, table in auxiliary_tables.items():
        path = output_dir / filename
        table.to_csv(path, index=False, encoding="utf-8-sig")
        paths[filename] = path
    return paths


def write_outputs(wide: pd.DataFrame, output_dir: Path, stem: str) -> tuple[Path, Path]:
    output_dir.mkdir(parents=True, exist_ok=True)
    csv_path = output_dir / f"{stem}.csv"
    xlsx_path = output_dir / f"{stem}.xlsx"
    wide.to_csv(csv_path, index=False, encoding="utf-8-sig")
    try:
        wide.to_excel(xlsx_path, index=False, engine="openpyxl")
    except PermissionError:
        fallback_path = output_dir / f"{stem}_cleaned.xlsx"
        wide.to_excel(fallback_path, index=False, engine="openpyxl")
        xlsx_path = fallback_path
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(xlsx_path)
        worksheet = workbook.active
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column_cells in worksheet.columns:
            header = str(column_cells[0].value or "")
            if header in {"body_raw", "body_clean", "body_without_metadata", "body_core", "retrieval_text", "rag_summary_text"}:
                width = 48
            elif header.startswith("evidence_") or header.endswith("_statement") or header.endswith("_text"):
                width = 36
            else:
                width = min(max(len(header) + 4, 14), 28)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        workbook.save(xlsx_path)
    except Exception as exc:  # pragma: no cover - formatting should not block data output
        print(f"Workbook formatting skipped: {exc}")
    return csv_path, xlsx_path


def count_body_core_date_residuals(wide: pd.DataFrame) -> int:
    count = 0
    for text in wide["body_core"].fillna("").astype(str):
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        if any(is_standalone_date_line(line) for line in lines):
            count += 1
    return count


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Build a deterministic qiaopi OCR wide table for RAG preprocessing.")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Input qiaopi OCR Excel workbook.")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="Worksheet name to process.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR, help="Directory for processed outputs.")
    parser.add_argument("--stem", default="qiaopi_213_wide_table", help="Output file stem without extension.")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.input.exists():
        raise FileNotFoundError(f"Input workbook not found: {args.input}")
    source = pd.read_excel(args.input, sheet_name=args.sheet)
    wide = build_wide_table(source)
    auxiliary_tables = build_auxiliary_tables(wide)
    auxiliary_paths = write_auxiliary_outputs(auxiliary_tables, args.output_dir)
    csv_path, xlsx_path = write_outputs(wide, args.output_dir, args.stem)
    print(f"Processed rows: {len(wide)}")
    print(f"Processed columns: {len(wide.columns)}")
    print("has_full_text distribution:")
    print(wide["has_full_text"].value_counts(dropna=False).sort_index().to_string())
    print("text_quality_level distribution:")
    print(wide["text_quality_level"].value_counts(dropna=False).to_string())
    print(f"review_needed count: {int(wide['review_needed'].sum())}")
    print("main_intent distribution:")
    print(wide["main_intent"].fillna("").value_counts(dropna=False).to_string())
    print(f"opening_salutation non-empty count: {int(wide['opening_salutation'].fillna('').astype(str).ne('').sum())}")
    print(f"signature non-empty count: {int(wide['signature'].fillna('').astype(str).ne('').sum())}")
    print(f"body_core date residual count: {count_body_core_date_residuals(wide)}")
    print("new field non-empty counts:")
    print(wide[NEW_WIDE_COLUMNS].fillna("").astype(str).ne("").sum().to_string())
    print(f"amount mention count: {len(auxiliary_tables['qiaopi_amount_mentions.csv'])}")
    print(f"place mention count: {len(auxiliary_tables['qiaopi_place_mentions.csv'])}")
    print(f"entity mention count: {len(auxiliary_tables['qiaopi_entity_mentions.csv'])}")
    print(f"evidence span count: {len(auxiliary_tables['qiaopi_evidence_spans.csv'])}")
    print(f"CSV output: {csv_path}")
    print(f"XLSX output: {xlsx_path}")
    for filename, path in auxiliary_paths.items():
        print(f"{filename}: {path}")


if __name__ == "__main__":
    main()
